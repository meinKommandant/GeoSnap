import logging
from pathlib import Path
from typing import List, Optional, Dict

import openpyxl

from .models import PhotoMetadata, GPSCoordinates

logger = logging.getLogger(__name__)


class ExcelImporter:
    """Importa metadatos de fotos desde un Excel (Source of Truth).

    Lee encabezados dinámicamente (case-insensitive) y soporta columnas con texto que
    contenga: "Nº", "Archivo", "DESCRIPCIÓN", "Fecha", "Latitud", "Longitud", "Altitud".
    """

    HEADER_KEYS = {
        # Removed "no" and "n" to avoid matching "Longitud" or "Filename"
        "num": ["nº", "numero", "n°", "id_foto"],
        "archivo": ["archivo", "file", "nombre", "filename"],
        "descripcion": ["descripción", "descripcion", "description", "notas"],
        "fecha": ["fecha", "date", "datetime", "timestamp"],
        "latitud": ["latitud", "lat", "latitude"],
        "longitud": ["longitud", "lon", "long", "lng", "longitude"],
        "altitud": ["altitud", "alt", "altitude", "elevacion"],
    }

    def parse_excel(self, excel_path: Path | str) -> List[PhotoMetadata]:
        excel_path = Path(excel_path)
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        ws = wb.active

        # Mapear encabezados de la fila 1
        header_map: Dict[str, int] = {}
        for cell in ws[1]:
            if cell.value is None:
                continue
            text = str(cell.value).strip().lower()
            if not text:
                continue
            for key, variants in self.HEADER_KEYS.items():
                if any(v in text for v in variants):
                    header_map[key] = cell.column  # 1-based index
        
        logger.info(f"Mapa de Encabezados detectado: {header_map}")

        # Validaciones críticas
        missing = [k for k in ("archivo", "latitud", "longitud") if k not in header_map]
        if missing:
            raise ValueError(
                f"Faltan columnas críticas en el Excel: {', '.join(missing)}. "
                "Asegúrate de incluir columnas para 'Archivo', 'Latitud' y 'Longitud'."
            )

        results: List[PhotoMetadata] = []

        for row_idx in range(2, ws.max_row + 1):
            def _val(col_key: str):
                col = header_map.get(col_key)
                return ws.cell(row=row_idx, column=col).value if col else None

            archivo = _val("archivo")
            if archivo is None:
                continue
            filename = str(archivo).strip()
            if not filename:
                continue

            # Latitud / Longitud obligatorias
            raw_lat = _val("latitud")
            raw_lon = _val("longitud")

            if raw_lat in (None, "") or raw_lon in (None, ""):
                logger.warning(f"Fila {row_idx}: Coordenadas faltantes. Se omite.")
                continue

            try:
                lat = self._to_float(raw_lat)
                lon = self._to_float(raw_lon)
            except ValueError:
                logger.warning(f"Fila {row_idx}: Coordenadas inválidas (Lat/Lon). Se omite.")
                continue

            # Altitud (opcional)
            alt_val = _val("altitud")
            try:
                altitude = self._to_float(alt_val) if alt_val not in (None, "") else 0.0
            except ValueError:
                altitude = 0.0

            # Fecha (opcional)
            fecha_cell = _val("fecha")
            timestamp = self._parse_datetime(fecha_cell)

            # Descripción (opcional)
            desc_cell = _val("descripcion")
            description = str(desc_cell).strip() if desc_cell not in (None, "") else ""

            # Sequence ID (Nº) - Source of Truth
            raw_num = _val("num")
            sequence_id = str(raw_num).strip() if raw_num is not None else None

            coords = GPSCoordinates(latitude=lat, longitude=lon, altitude=altitude)
            results.append(
                PhotoMetadata(
                    filename=filename,
                    filepath="",  # Se resolverá más tarde a partir del índice de fotos
                    timestamp=timestamp,
                    coordinates=coords,
                    description=description,
                    sequence_id=sequence_id,
                )
            )

        return results

    def _to_float(self, value) -> float:
        if isinstance(value, (int, float)):
            return float(value)
        s = str(value).strip().replace(",", ".")
        return float(s)

    def _parse_datetime(self, value):
        # openpyxl suele retornar datetime para celdas de fecha válidas
        from datetime import datetime, date
        if value is None or value == "":
            return None
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, datetime.min.time())
        try:
            # Intento de parseo de string común
            # 2024-01-31 12:34:56 o 31/01/2024 12:34:56
            txt = str(value).strip()
            for fmt in ("%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y"):
                try:
                    return datetime.strptime(txt, fmt)
                except Exception:
                    pass
        except Exception:
            pass
        # Si es inválido, devolver None (se intentará llenar luego)
        return None
